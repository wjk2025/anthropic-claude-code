"""Excel/XLSX export functionality for office space reports."""

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, NamedStyle
from openpyxl.utils import get_column_letter

from ..calculator.space_calculator import CalculationResult
from ..calculator.remote_work_analyzer import RemoteWorkScenario
from ..models.space_standards import DEFAULT_SPACE_STANDARDS


class ExcelExporter:
    """Exports space program calculations to Excel/XLSX format."""

    def __init__(self, output_dir: Optional[Path] = None):
        """Initialize exporter with optional output directory."""
        self.output_dir = output_dir or Path(".")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self._setup_styles()

    def _setup_styles(self):
        """Set up reusable styles."""
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        self.title_font = Font(bold=True, size=16, color="2C3E50")
        self.subtitle_font = Font(bold=True, size=12, color="34495E")
        self.section_font = Font(bold=True, size=11, color="2C3E50")

        self.number_font = Font(size=10)
        self.total_font = Font(bold=True, size=10)
        self.total_fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")

        thin_border = Side(style="thin", color="CCCCCC")
        self.cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    def _apply_header_style(self, cell):
        """Apply header styling to a cell."""
        cell.font = self.header_font
        cell.fill = self.header_fill
        cell.alignment = self.header_alignment
        cell.border = self.cell_border

    def _apply_data_style(self, cell, is_number: bool = False):
        """Apply data cell styling."""
        cell.font = self.number_font
        cell.border = self.cell_border
        if is_number:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="left")

    def _apply_total_style(self, cell, is_number: bool = False):
        """Apply total row styling."""
        cell.font = self.total_font
        cell.fill = self.total_fill
        cell.border = self.cell_border
        if is_number:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0"

    def _auto_column_width(self, ws, min_width: int = 10, max_width: int = 50):
        """Auto-adjust column widths based on content."""
        for column_cells in ws.columns:
            length = max(
                min(len(str(cell.value or "")) + 2, max_width)
                for cell in column_cells
            )
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(length, min_width)

    def export(
        self,
        result: CalculationResult,
        filename: Optional[str] = None,
        include_remote_analysis: bool = False,
        remote_scenarios: Optional[list[RemoteWorkScenario]] = None,
    ) -> Path:
        """Export calculation results to Excel."""
        project = result.project
        summary = result.summary

        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_name = "".join(c if c.isalnum() else "_" for c in project.company_name)
            filename = f"space_program_{safe_name}_{timestamp}.xlsx"

        output_path = self.output_dir / filename
        wb = Workbook()

        # Summary Sheet
        ws_summary = wb.active
        ws_summary.title = "Summary"
        self._create_summary_sheet(ws_summary, result)

        # Department Details Sheet
        ws_dept = wb.create_sheet("Departments")
        self._create_department_sheet(ws_dept, result)

        # Space Breakdown Sheet
        ws_spaces = wb.create_sheet("Space Breakdown")
        self._create_space_breakdown_sheet(ws_spaces, result)

        # Standards Reference Sheet
        ws_standards = wb.create_sheet("Space Standards")
        self._create_standards_sheet(ws_standards, project)

        # Remote Work Analysis Sheet
        if include_remote_analysis and remote_scenarios:
            ws_remote = wb.create_sheet("Remote Work Analysis")
            self._create_remote_analysis_sheet(ws_remote, remote_scenarios)

        wb.save(output_path)
        return output_path

    def _create_summary_sheet(self, ws, result: CalculationResult):
        """Create the summary sheet."""
        project = result.project
        summary = result.summary
        row = 1

        # Title
        ws.cell(row=row, column=1, value="OFFICE SPACE PROGRAM").font = self.title_font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 2

        # Project info
        info = [
            ("Company:", project.company_name),
            ("Project:", project.project_name or "-"),
            ("Location:", project.location or "-"),
            ("Prepared:", datetime.now().strftime("%B %d, %Y")),
            ("Prepared By:", project.prepared_by or "-"),
        ]

        for label, value in info:
            ws.cell(row=row, column=1, value=label).font = self.section_font
            ws.cell(row=row, column=2, value=value)
            row += 1

        row += 2

        # Executive Summary
        ws.cell(row=row, column=1, value="EXECUTIVE SUMMARY").font = self.subtitle_font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1

        headers = ["Metric", "Value", "Unit", "Notes"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        summary_data = [
            ("Total Headcount", summary.total_headcount, "people", ""),
            ("Effective On-Site", summary.effective_onsite, "people", "Accounting for remote work"),
            ("", "", "", ""),
            ("Net Usable Area", summary.net_usable_sf, "SF", "Workspace + Meeting + Support"),
            ("Circulation Factor", summary.circulation_factor * 100, "%", "Industry standard: 25-35%"),
            ("Circulation Area", summary.circulation_sf, "SF", ""),
            ("Total Usable Area", summary.usable_sf, "SF", "Net + Circulation"),
            ("", "", "", ""),
        ]

        if summary.loss_factor > 0:
            summary_data.extend([
                ("Loss Factor", summary.loss_factor * 100, "%", "Building efficiency factor"),
                ("Total Rentable Area", summary.rentable_sf, "SF", "Usable × (1 + Loss Factor)"),
                ("", "", "", ""),
            ])

        summary_data.extend([
            ("Usable SF per Person", summary.sf_per_person, "SF/person", ""),
        ])
        if summary.loss_factor > 0:
            summary_data.append(
                ("Rentable SF per Person", summary.sf_per_person_rentable, "SF/person", "")
            )

        for metric, value, unit, notes in summary_data:
            if metric:
                ws.cell(row=row, column=1, value=metric)
                cell = ws.cell(row=row, column=2, value=value)
                self._apply_data_style(cell, is_number=isinstance(value, (int, float)))
                ws.cell(row=row, column=3, value=unit)
                ws.cell(row=row, column=4, value=notes)
            row += 1

        row += 2

        # Area Breakdown
        ws.cell(row=row, column=1, value="AREA BREAKDOWN").font = self.subtitle_font
        row += 1

        headers = ["Category", "Square Feet", "% of Net"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        areas = [
            ("Workspaces", summary.total_workspace_sf),
            ("Meeting Spaces", summary.total_meeting_sf),
            ("Support Spaces", summary.total_support_sf),
        ]

        for category, sf in areas:
            ws.cell(row=row, column=1, value=category)
            cell = ws.cell(row=row, column=2, value=sf)
            self._apply_data_style(cell, is_number=True)
            pct = (sf / summary.net_usable_sf * 100) if summary.net_usable_sf > 0 else 0
            ws.cell(row=row, column=3, value=f"{pct:.1f}%")
            row += 1

        # Total row
        ws.cell(row=row, column=1, value="NET USABLE TOTAL")
        self._apply_total_style(ws.cell(row=row, column=1))
        cell = ws.cell(row=row, column=2, value=summary.net_usable_sf)
        self._apply_total_style(cell, is_number=True)
        self._apply_total_style(ws.cell(row=row, column=3, value="100%"))

        self._auto_column_width(ws)

    def _create_department_sheet(self, ws, result: CalculationResult):
        """Create the department details sheet."""
        project = result.project
        row = 1

        ws.cell(row=row, column=1, value="DEPARTMENT BREAKDOWN").font = self.title_font
        row += 2

        # Summary table
        headers = ["Department", "Headcount", "Workspace SF", "SF/Person"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        for dept in project.departments:
            dept_sf = result.department_workspace_sf.get(dept.name, 0)
            sf_per = dept_sf / dept.total_headcount if dept.total_headcount > 0 else 0

            ws.cell(row=row, column=1, value=dept.name)
            cell = ws.cell(row=row, column=2, value=dept.total_headcount)
            self._apply_data_style(cell, is_number=True)
            cell = ws.cell(row=row, column=3, value=dept_sf)
            self._apply_data_style(cell, is_number=True)
            cell = ws.cell(row=row, column=4, value=sf_per)
            self._apply_data_style(cell, is_number=True)
            row += 1

        # Total
        total_hc = result.summary.total_headcount
        total_sf = result.summary.total_workspace_sf
        avg_sf = total_sf / total_hc if total_hc > 0 else 0

        self._apply_total_style(ws.cell(row=row, column=1, value="TOTAL"))
        self._apply_total_style(ws.cell(row=row, column=2, value=total_hc), is_number=True)
        self._apply_total_style(ws.cell(row=row, column=3, value=total_sf), is_number=True)
        self._apply_total_style(ws.cell(row=row, column=4, value=avg_sf), is_number=True)
        row += 3

        # Staff detail by department
        ws.cell(row=row, column=1, value="STAFF DETAIL BY DEPARTMENT").font = self.subtitle_font
        row += 2

        for dept in project.departments:
            ws.cell(row=row, column=1, value=dept.name).font = self.section_font
            row += 1

            headers = ["Role", "Count", "Workspace Type", "Remote %", "SF/Unit", "Total SF"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                self._apply_header_style(cell)
            row += 1

            for staff in dept.staff:
                size = project.get_space_size(staff.space_type)
                total = staff.count * size

                ws.cell(row=row, column=1, value=staff.role)
                self._apply_data_style(ws.cell(row=row, column=2, value=staff.count), is_number=True)
                ws.cell(row=row, column=3, value=staff.space_type.replace("_", " ").title())
                ws.cell(row=row, column=4, value=f"{staff.remote_percentage:.0f}%")
                self._apply_data_style(ws.cell(row=row, column=5, value=size), is_number=True)
                self._apply_data_style(ws.cell(row=row, column=6, value=total), is_number=True)
                row += 1

            row += 1

        self._auto_column_width(ws)

    def _create_space_breakdown_sheet(self, ws, result: CalculationResult):
        """Create the space breakdown sheet."""
        summary = result.summary
        project = result.project
        row = 1

        ws.cell(row=row, column=1, value="SPACE BREAKDOWN").font = self.title_font
        row += 2

        # Workspace types
        if result.workspace_counts:
            ws.cell(row=row, column=1, value="WORKSPACES").font = self.subtitle_font
            row += 1

            headers = ["Space Type", "Count", "SF/Unit", "Total SF"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                self._apply_header_style(cell)
            row += 1

            total_count = 0
            total_sf = 0
            for space_type, count in result.workspace_counts.items():
                if count > 0:
                    size = project.get_space_size(space_type)
                    sf = count * size
                    total_count += count
                    total_sf += sf

                    ws.cell(row=row, column=1, value=space_type.replace("_", " ").title())
                    self._apply_data_style(ws.cell(row=row, column=2, value=count), is_number=True)
                    self._apply_data_style(ws.cell(row=row, column=3, value=size), is_number=True)
                    self._apply_data_style(ws.cell(row=row, column=4, value=sf), is_number=True)
                    row += 1

            self._apply_total_style(ws.cell(row=row, column=1, value="SUBTOTAL"))
            self._apply_total_style(ws.cell(row=row, column=2, value=total_count), is_number=True)
            self._apply_total_style(ws.cell(row=row, column=3, value=""))
            self._apply_total_style(ws.cell(row=row, column=4, value=total_sf), is_number=True)
            row += 2

        # Meeting spaces
        if summary.meeting_breakdown:
            ws.cell(row=row, column=1, value="MEETING SPACES").font = self.subtitle_font
            row += 1

            headers = ["Space Type", "Quantity", "Total SF"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                self._apply_header_style(cell)
            row += 1

            for name, sf in summary.meeting_breakdown.items():
                ws.cell(row=row, column=1, value=name)
                qty = next((v for k, v in result.meeting_counts.items()
                           if DEFAULT_SPACE_STANDARDS[k].name == name), 1)
                self._apply_data_style(ws.cell(row=row, column=2, value=qty), is_number=True)
                self._apply_data_style(ws.cell(row=row, column=3, value=sf), is_number=True)
                row += 1

            self._apply_total_style(ws.cell(row=row, column=1, value="SUBTOTAL"))
            self._apply_total_style(ws.cell(row=row, column=2, value=""))
            self._apply_total_style(ws.cell(row=row, column=3, value=summary.total_meeting_sf), is_number=True)
            row += 2

        # Support spaces
        if summary.support_breakdown:
            ws.cell(row=row, column=1, value="SUPPORT SPACES").font = self.subtitle_font
            row += 1

            headers = ["Space Type", "Total SF"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                self._apply_header_style(cell)
            row += 1

            for name, sf in summary.support_breakdown.items():
                ws.cell(row=row, column=1, value=name)
                self._apply_data_style(ws.cell(row=row, column=2, value=sf), is_number=True)
                row += 1

            self._apply_total_style(ws.cell(row=row, column=1, value="SUBTOTAL"))
            self._apply_total_style(ws.cell(row=row, column=2, value=summary.total_support_sf), is_number=True)

        self._auto_column_width(ws)

    def _create_standards_sheet(self, ws, project):
        """Create the space standards reference sheet."""
        row = 1

        ws.cell(row=row, column=1, value="SPACE STANDARDS REFERENCE").font = self.title_font
        row += 2

        ws.cell(row=row, column=1, value="Standard sizes used in this program (can be customized per project)").font = Font(italic=True)
        row += 2

        headers = ["Space Type", "Standard SF", "Custom SF", "Description", "Capacity"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        for space_type, standard in DEFAULT_SPACE_STANDARDS.items():
            custom_sf = project.custom_standards.get(space_type)

            ws.cell(row=row, column=1, value=standard.name)
            self._apply_data_style(ws.cell(row=row, column=2, value=standard.square_feet), is_number=True)
            if custom_sf:
                self._apply_data_style(ws.cell(row=row, column=3, value=custom_sf), is_number=True)
            else:
                ws.cell(row=row, column=3, value="-")
            ws.cell(row=row, column=4, value=standard.description)
            ws.cell(row=row, column=5, value=standard.capacity or "-")
            row += 1

        self._auto_column_width(ws)

    def _create_remote_analysis_sheet(self, ws, scenarios: list[RemoteWorkScenario]):
        """Create the remote work analysis sheet."""
        row = 1

        ws.cell(row=row, column=1, value="REMOTE WORK IMPACT ANALYSIS").font = self.title_font
        row += 2

        ws.cell(row=row, column=1, value="Analysis of space requirements under different remote work scenarios").font = Font(italic=True)
        row += 2

        headers = ["Scenario", "Remote %", "Days In Office", "Baseline SF", "Adjusted SF", "Reduction SF", "Reduction %", "Est. Annual Savings"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        for scenario in scenarios:
            reduction_pct = (
                (scenario.baseline_usable_sf - scenario.adjusted_usable_sf)
                / scenario.baseline_usable_sf * 100
                if scenario.baseline_usable_sf > 0 else 0
            )

            ws.cell(row=row, column=1, value=scenario.name)
            ws.cell(row=row, column=2, value=f"{scenario.remote_percentage:.0f}%")
            ws.cell(row=row, column=3, value=f"{scenario.days_in_office_per_week:.1f}")
            self._apply_data_style(ws.cell(row=row, column=4, value=scenario.baseline_usable_sf), is_number=True)
            self._apply_data_style(ws.cell(row=row, column=5, value=scenario.adjusted_usable_sf), is_number=True)
            self._apply_data_style(ws.cell(row=row, column=6, value=scenario.sf_reduction), is_number=True)
            ws.cell(row=row, column=7, value=f"{reduction_pct:.1f}%")
            cell = ws.cell(row=row, column=8, value=scenario.cost_savings_estimate)
            cell.number_format = '"$"#,##0'
            row += 1

        row += 2

        # Scenario descriptions
        ws.cell(row=row, column=1, value="SCENARIO DESCRIPTIONS").font = self.subtitle_font
        row += 1

        for scenario in scenarios:
            ws.cell(row=row, column=1, value=scenario.name).font = self.section_font
            ws.cell(row=row, column=2, value=scenario.description)
            row += 1

        self._auto_column_width(ws)
